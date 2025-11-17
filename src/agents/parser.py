"""Agente Parser - Extrai e normaliza perguntas de questionários."""
from typing import List, Dict, Any, Optional
from pathlib import Path
from langchain_core.language_models import BaseChatModel
from langchain_core.messages import HumanMessage, SystemMessage
from src.workflows.schema import ParsedQuestion
from src.utils.metrics import (
    parser_extractions_total,
    parser_extraction_duration_seconds,
    parser_questions_normalized_total
)
import pdfplumber
import PyPDF2
from docx import Document
import pandas as pd
import json
import time
import logging

logger = logging.getLogger(__name__)

# Tentar importar OCR (opcional)
try:
    from PIL import Image
    import pytesseract
    OCR_AVAILABLE = True
except ImportError:
    OCR_AVAILABLE = False
    logger.warning("OCR não disponível (PIL/pytesseract não instalado)")


class ParserAgent:
    """Agente Parser para extração e normalização de perguntas."""
    
    def __init__(self, llm: BaseChatModel):
        """Inicializar Agente Parser."""
        self.llm = llm
        self._setup_system_prompt()
    
    def _setup_system_prompt(self) -> None:
        """Configurar prompt do sistema."""
        self.system_prompt = """Você é um Agente Parser especializado em extrair e normalizar perguntas de questionários.
Sua função é:
1. Extrair perguntas de diferentes formatos
2. Normalizar em formato canônico
3. Identificar categoria e requisitos
4. Mapear para perguntas similares do histórico

Sempre responda em formato JSON estruturado."""
    
    def extract_from_file(self, file_path: str, use_ocr: bool = False) -> str:
        """Extrair texto de arquivo baseado na extensão."""
        file_path_obj = Path(file_path)
        extension = file_path_obj.suffix.lower()
        
        start_time = time.time()
        
        try:
            if extension == '.pdf':
                text = self.extract_from_pdf(file_path, use_ocr=use_ocr)
                file_type = 'pdf'
            elif extension in ['.docx', '.doc']:
                text = self.extract_from_docx(file_path)
                file_type = 'docx'
            elif extension in ['.xlsx', '.xls']:
                text = self.extract_from_excel(file_path)
                file_type = 'excel'
            elif extension == '.csv':
                text = self.extract_from_csv(file_path)
                file_type = 'csv'
            else:
                raise ValueError(f"Formato de arquivo não suportado: {extension}")
            
            duration = time.time() - start_time
            
            # Métricas
            parser_extractions_total.labels(
                file_type=file_type,
                status='success'
            ).inc()
            parser_extraction_duration_seconds.labels(
                file_type=file_type
            ).observe(duration)
            
            logger.info(f"Texto extraído de {file_path} ({file_type}): {len(text)} caracteres")
            return text
            
        except Exception as e:
            duration = time.time() - start_time
            file_type = extension.lstrip('.')
            
            parser_extractions_total.labels(
                file_type=file_type,
                status='error'
            ).inc()
            parser_extraction_duration_seconds.labels(
                file_type=file_type
            ).observe(duration)
            
            logger.error(f"Erro ao extrair texto de {file_path}: {e}")
            raise
    
    def extract_from_pdf(self, file_path: str, use_ocr: bool = False) -> str:
        """Extrair texto de PDF."""
        text_content = []
        
        try:
            # Tentar extração direta primeiro com pdfplumber
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    text = page.extract_text()
                    if text and text.strip():
                        text_content.append(text)
            
            # Se texto vazio e OCR habilitado, usar OCR
            if not text_content and use_ocr:
                if OCR_AVAILABLE:
                    logger.info("Tentando OCR para PDF escaneado")
                    text_content = self._extract_with_ocr(file_path)
                else:
                    logger.warning("OCR solicitado mas não disponível")
            
            if not text_content:
                # Fallback para PyPDF2
                logger.debug("Tentando PyPDF2 como fallback")
                with open(file_path, 'rb') as file:
                    pdf_reader = PyPDF2.PdfReader(file)
                    for page in pdf_reader.pages:
                        text = page.extract_text()
                        if text and text.strip():
                            text_content.append(text)
            
            return "\n".join(text_content) if text_content else ""
            
        except Exception as e:
            logger.error(f"Erro ao extrair PDF: {e}")
            raise
    
    def _extract_with_ocr(self, file_path: str) -> List[str]:
        """Extrair texto usando OCR."""
        if not OCR_AVAILABLE:
            logger.warning("OCR não disponível")
            return []
        
        try:
            from pdf2image import convert_from_path
            images = convert_from_path(file_path)
            text_pages = []
            
            for image in images:
                text = pytesseract.image_to_string(image, lang='por')
                if text.strip():
                    text_pages.append(text)
            
            logger.info(f"OCR extraiu {len(text_pages)} páginas")
            return text_pages
            
        except Exception as e:
            logger.error(f"Erro no OCR: {e}")
            return []
    
    def extract_from_docx(self, file_path: str) -> str:
        """Extrair texto de DOCX."""
        try:
            doc = Document(file_path)
            paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
            
            # Também extrair tabelas
            tables_text = []
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        tables_text.append(" | ".join(row_text))
            
            all_text = paragraphs + tables_text
            return "\n".join(all_text)
            
        except Exception as e:
            logger.error(f"Erro ao extrair DOCX: {e}")
            raise
    
    def extract_from_excel(self, file_path: str) -> str:
        """Extrair texto de Excel."""
        try:
            # Ler todas as planilhas
            excel_file = pd.ExcelFile(file_path)
            text_parts = []
            
            for sheet_name in excel_file.sheet_names:
                df = pd.read_excel(excel_file, sheet_name=sheet_name)
                text_parts.append(f"Sheet: {sheet_name}")
                text_parts.append(df.to_string(index=False))
                text_parts.append("")  # Linha em branco entre sheets
            
            return "\n".join(text_parts)
            
        except Exception as e:
            logger.error(f"Erro ao extrair Excel: {e}")
            raise
    
    def extract_from_csv(self, file_path: str) -> str:
        """Extrair texto de CSV."""
        try:
            df = pd.read_csv(file_path)
            return df.to_string(index=False)
            
        except Exception as e:
            logger.error(f"Erro ao extrair CSV: {e}")
            raise
    
    @observe(name="parser_normalize_questions")
    def normalize_questions(self, raw_text: str) -> List[ParsedQuestion]:
        """Normalizar perguntas em formato canônico."""
        logger.info("Iniciando normalização de perguntas")
        logger.debug(f"Texto para normalizar: {raw_text[:200]}...")
        
        prompt = f"""Extraia e normalize as seguintes perguntas em formato JSON:

Para cada pergunta, identifique:
- qid: Identificador único (ex: Q001, Q002, Q003)
- category: Categoria (técnico, segurança, compliance, jurídico, comercial, geral)
- question_text: Texto da pergunta normalizado e limpo
- requirements: Lista de requisitos mencionados na pergunta
- expected_format: Formato esperado da resposta (texto, número, sim/não, lista, tabela, etc.)

Texto: {raw_text}

Responda APENAS com um JSON array válido, sem texto adicional:
[{{"qid": "Q001", "category": "técnico", "question_text": "...", "requirements": ["..."], "expected_format": "texto"}}]"""
        
        messages = [
            SystemMessage(content=self.system_prompt),
            HumanMessage(content=prompt)
        ]
        
        try:
            response = self.llm.invoke(messages)
            content = response.content if hasattr(response, 'content') else str(response)
            
            # Extrair JSON
            json_start = content.find('[')
            json_end = content.rfind(']') + 1
            
            if json_start >= 0 and json_end > json_start:
                json_str = content[json_start:json_end]
                questions_data = json.loads(json_str)
                
                parsed_questions = []
                for q_data in questions_data:
                    try:
                        # Garantir que todos os campos obrigatórios existem
                        if 'qid' not in q_data:
                            q_data['qid'] = f"Q{len(parsed_questions) + 1:03d}"
                        if 'requirements' not in q_data:
                            q_data['requirements'] = []
                        if 'expected_format' not in q_data:
                            q_data['expected_format'] = 'texto'
                        
                        parsed_question = ParsedQuestion(**q_data)
                        parsed_questions.append(parsed_question)
                        
                        # Métricas
                        parser_questions_normalized_total.labels(
                            category=parsed_question.category
                        ).inc()
                        
                    except Exception as e:
                        logger.warning(f"Erro ao criar ParsedQuestion: {e}, dados: {q_data}")
                        continue
                
                logger.info(f"Normalizadas {len(parsed_questions)} perguntas")
                return parsed_questions
            else:
                raise ValueError("JSON array não encontrado na resposta")
                
        except json.JSONDecodeError as e:
            logger.error(f"Erro ao fazer parse do JSON: {e}")
            logger.debug(f"Resposta recebida: {content[:500]}")
            return []
        except Exception as e:
            logger.error(f"Erro ao normalizar perguntas: {e}")
            return []
    
    def map_to_historical(
        self, 
        questions: List[ParsedQuestion],
        historical_questions: Optional[List[Dict[str, Any]]] = None
    ) -> List[Dict[str, Any]]:
        """Mapear perguntas para histórico usando similaridade semântica."""
        logger.info(f"Mapeando {len(questions)} perguntas para histórico")
        
        # TODO: Implementar busca semântica no vector store quando disponível
        # Por enquanto, retornar estrutura básica
        mapped = []
        for q in questions:
            mapped.append({
                "qid": q.qid,
                "question": q.question_text,
                "category": q.category,
                "similar_questions": q.similar_questions,
                "mapping_confidence": 0.0  # Será calculado com vector search
            })
        
        logger.debug(f"Mapeamento concluído para {len(mapped)} perguntas")
        return mapped
    
    def process(self, input_data: str, file_path: Optional[str] = None) -> List[ParsedQuestion]:
        """Processar input (texto ou arquivo) e retornar perguntas normalizadas."""
        logger.info("Iniciando processamento do Parser")
        
        # Se tem arquivo, extrair texto primeiro
        if file_path:
            raw_text = self.extract_from_file(file_path)
        else:
            raw_text = input_data
        
        # Normalizar perguntas
        parsed_questions = self.normalize_questions(raw_text)
        
        # Mapear para histórico
        if parsed_questions:
            mapped = self.map_to_historical(parsed_questions)
            logger.info(f"Processamento concluído: {len(parsed_questions)} perguntas")
        
        return parsed_questions
